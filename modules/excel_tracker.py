"""
Excel Tracker: Records content details and video links for each pipeline run.
Maintains a content_log.xlsx file for tracking all generated content.
"""

import os
import datetime
from config import Config

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExcelTracker:
    """Tracks all pipeline runs in an Excel spreadsheet."""

    HEADERS = [
        "Sr. No",
        "Run ID",
        "Date & Time",
        "Topic Title",
        "Category",
        "Premise",
        "Script Hook",
        "Characters",
        "Video Path",
        "YouTube Link",
        "Status",
    ]

    def __init__(self, filepath=None):
        if filepath is None:
            self.filepath = os.path.join(Config.BASE_DIR, "content_log.xlsx")
        else:
            self.filepath = filepath

    def _create_workbook(self):
        """Create a new workbook with formatted headers."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Content Log"

        # Style the header row
        header_font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_idx, header in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Set column widths
        widths = [8, 22, 20, 25, 18, 40, 30, 25, 40, 45, 12]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

        # Freeze the header row
        ws.freeze_panes = "A2"

        return wb

    def log_run(self, run_id, topic=None, script=None, video_path=None,
                upload_result=None, status="Completed"):
        """Log a pipeline run to the Excel file.

        Args:
            run_id: The run folder name (e.g., 'run_20260218_153533').
            topic: dict with topic info (title, category, premise).
            script: dict with script info (hook, characters).
            video_path: Path to the final video file.
            upload_result: dict with YouTube upload result (url, video_id).
            status: Run status ('Completed', 'Failed', 'Partial').
        """
        if not OPENPYXL_AVAILABLE:
            print("   ⚠️  openpyxl not installed. Cannot log to Excel.")
            print("   Run: pip install openpyxl")
            return

        # Load or create workbook
        if os.path.exists(self.filepath):
            wb = load_workbook(self.filepath)
            ws = wb.active
        else:
            wb = self._create_workbook()
            ws = wb.active

        # Determine row number (Sr. No)
        next_row = ws.max_row + 1
        sr_no = next_row - 1  # minus header

        # Extract data from inputs
        topic = topic or {}
        script = script or {}
        upload_result = upload_result or {}

        title = topic.get("title", "N/A")
        category = topic.get("category", "N/A")
        premise = topic.get("premise", "N/A")
        hook = script.get("hook", "N/A")
        characters = ", ".join(script.get("characters", []))
        youtube_link = upload_result.get("url", "Not uploaded")
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # Check if video exists
        if video_path and os.path.exists(video_path):
            video_status = video_path
        else:
            video_status = video_path or "Not generated"

        row_data = [
            sr_no,
            run_id,
            now,
            title,
            category,
            premise,
            hook,
            characters,
            video_status,
            youtube_link,
            status,
        ]

        # Style for data rows
        data_align = Alignment(vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Color code status
        status_colors = {
            "Completed": "C6EFCE",
            "Failed": "FFC7CE",
            "Partial": "FFEB9C",
        }

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=next_row, column=col_idx, value=value)
            cell.alignment = data_align
            cell.border = thin_border

            # Apply status color
            if col_idx == len(self.HEADERS):  # Status column
                fill_color = status_colors.get(status, "FFFFFF")
                cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        wb.save(self.filepath)
        print(f"   📊 Logged to Excel: {self.filepath}")
        return self.filepath
